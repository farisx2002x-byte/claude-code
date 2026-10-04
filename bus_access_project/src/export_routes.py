"""طبقات المسارات + KML المسارات (بدون أسماء) + جداول_السائقين.xlsx (داخلي، فيه أسماء)."""
import html

import geopandas as gpd
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from shapely.geometry import LineString

import config as C
from src.export_gis import lonlat


def gis_layers(routes, route_stops, geoms):
    rows, geo = [], []
    for r in routes.itertuples():
        g = geoms.get(r.route)
        if g is None or len(g) < 2:
            continue
        rows.append(dict(route=r.route, school=r.school, vehicle=C.VEHICLES[r.vk]["name"], stops=r.stops,
                         students=r.students, seats=r.seats, duration_min=round(r.duration_min, 1),
                         longest_ride_min=round(r.longest_ride_min, 1), length_km=round(r.length_km, 2),
                         start_time=r.start_time))
        geo.append(LineString(g))
    rg = gpd.GeoDataFrame(rows, geometry=geo, crs=C.CRS_UTM) if rows else None
    sg = gpd.GeoDataFrame(route_stops.drop(columns=["x", "y"]),
                          geometry=gpd.points_from_xy(route_stops.x, route_stops.y), crs=C.CRS_UTM) if len(route_stops) else None
    return rg, sg


def write_routes_kml(path, routes, route_stops, geoms):
    """مجلد لكل مدرسة ← مجلد لكل مسار، فيه الخط والوقفات المرقمة. بدون أسماء طلاب."""
    out = ['<?xml version="1.0" encoding="UTF-8"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document><name>مسارات الباصات</name>']
    for school, rr in routes.groupby("school"):
        out.append(f"<Folder><name>{html.escape(school)}</name>")
        for r in rr.itertuples():
            out.append(f"<Folder><name>{html.escape(r.route)}</name>")
            g = geoms.get(r.route)
            if g is not None and len(g) > 1:
                lon, lat = lonlat(g[:, 0], g[:, 1])
                co = " ".join(f"{a:.6f},{b:.6f},0" for a, b in zip(lon, lat))
                out.append(f"<Placemark><name>{html.escape(r.route)}</name><LineString><coordinates>{co}</coordinates></LineString></Placemark>")
            st = route_stops[route_stops.route == r.route].sort_values("order")
            lon, lat = lonlat(st.x, st.y)
            for k, s in enumerate(st.itertuples()):
                out.append(f"<Placemark><name>{s.order}</name><description>{s.time} | {s.students} طالب</description>"
                           f"<Point><coordinates>{lon[k]:.6f},{lat[k]:.6f},0</coordinates></Point></Placemark>")
            out.append("</Folder>")
        out.append("</Folder>")
    out.append("</Document></kml>")
    open(path, "w", encoding="utf-8").write("".join(out))


def write_driver_sheets(path, routes, route_stops, assign, students, no_route_idx):
    """ملف داخلي للسائقين: فيه أسماء الطلاب وأرقام الوقفات."""
    wb = Workbook()
    wb.remove(wb.active)
    names = students.set_index("idx")["name"].to_dict()

    def sheet(title, df):
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, "item") else v) for v in row])
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
        for row in ws.iter_rows():
            for c in row:
                c.alignment = Alignment(readingOrder=2, wrap_text=True, vertical="top")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 22
        ws.freeze_panes = "A2"
    r = routes.rename(columns={"route": "المسار", "school": "المدرسة", "stops": "الوقفات", "students": "الطلاب",
                               "seats": "المقاعد", "duration_min": "المدة (د)", "longest_ride_min": "أطول رحلة (د)",
                               "length_km": "الطول (كم)", "start_time": "الانطلاق"})
    r["المركبة"] = routes["vk"].map(lambda v: C.VEHICLES[v]["name"])
    sheet("ملخص المسارات", r[["المسار", "المدرسة", "المركبة", "الوقفات", "الطلاب", "المقاعد", "الانطلاق", "المدة (د)", "أطول رحلة (د)", "الطول (كم)"]].round(1))
    rows = []
    st_assign = assign.groupby("route")
    ass_by_route = {k: v for k, v in st_assign}
    for rs in route_stops.sort_values(["route", "order"]).itertuples():
        # طلاب الوقفة: نطابق بوقت الركوب نفسه داخل المسار (الوقفات المجمعة تشترك في الوقت)
        sub = ass_by_route.get(rs.route)
        mem = sub[(sub.board_time == rs.time) & (abs(sub.ride_min - rs.ride_min) < 1e-6)] if sub is not None else []
        lon, lat = lonlat([rs.x], [rs.y])
        rows.append({"المسار": rs.route, "الترتيب": rs.order, "الوقت": rs.time, "العدد": rs.students,
                     "الأسماء": "، ".join(names.get(i, "") for i in (mem["idx"] if len(mem) else [])),
                     "الموقع": f"https://www.google.com/maps?q={lat[0]:.6f},{lon[0]:.6f}"})
    sheet("الوقفات والطلاب", pd.DataFrame(rows))
    nr = students[students["idx"].isin(no_route_idx)]
    sheet("طلاب بدون مسار", pd.DataFrame({"الاسم": nr["name"], "المدرسة": nr["school"], "الحي": nr.get("district", "")}))
    sheet("ملاحظات", pd.DataFrame({"ملاحظة": [
        "هذا الملف داخلي ويحتوي أسماء قُصّر، لا يُرسل خارج الجهة.",
        "الأوقات محسوبة بالرجوع من وصول المدرسة " + C.SCHOOL_ARRIVAL + " وبسرعات تقديرية، وتحتاج معايرة مع المشغّل.",
        "الوقفات تُدمج إذا كانت أقرب من " + str(C.STOP_MERGE_M) + " م لنفس الفئة."]}))
    wb.save(path)
