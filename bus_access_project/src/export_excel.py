"""تقرير Excel بالعربي (RTL): كل الأوراق من اليمين لليسار، خط Arial، ومعادلات حيّة."""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as CL

import config as C
from src.export_gis import lonlat

FONT = "Arial"
FILL = {"سهل": "C6EFCE", "متوسط": "FFEB9C", "صعب": "FFC7CE"}
HDR = PatternFill("solid", fgColor="1F3864")


def _sheet(wb, title):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    return ws


def _put_df(ws, df, start_row=1, widths=None):
    cols = list(df.columns)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=start_row, column=j, value=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, float) and (np.isnan(v) or np.isinf(v)):
                v = None
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (np.floating,)):
                v = float(v)
            elif isinstance(v, (np.bool_,)):
                v = bool(v)
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FONT)
            c.alignment = Alignment(readingOrder=2, vertical="top")
    for j, c in enumerate(cols, 1):
        w = (widths or {}).get(c) or min(max(len(str(c)) + 4, 12), 40)
        ws.column_dimensions[CL(j)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def _level_cf(ws, rng):
    for lv, color in FILL.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{lv}"'],
                                                      fill=PatternFill("solid", bgColor=color, fgColor=color)))


def _col(df_cols, name):
    return CL(list(df_cols).index(name) + 1)


METHOD = [
    "المنهجية (الإصدار " + C.VERSION + ")",
    "عرض الشارع المقاس = حرم الشارع بين حدود القطع (مو الإسفلت)، ودقة حدود الأمانة تقريباً ±1–2 م.",
    "كل مركبة تُحلل بحدودها (أقل حرم، الالتفاف الحاد، مساحة الدوران) على شبكة الشوارع المقاسة.",
    "المشي: على شبكة الشوارع، وأول 30 م مجانية، ويُضرب ×1.4 للابتدائي والطفولة المبكرة.",
    "المستوى: 0–2 سهل، 3–5 متوسط، 6+ صعب. وصعب مباشرة لو ما فيه وصول أو المشي بعد المعامل > 250 م.",
    "المركبة الموصى بها: أكبر مركبة مستواها سهل، وإلا أكبر مركبة مستواها متوسط، وإلا نقطة تجميع.",
    "المسارات: OR-Tools بأسطول مختلط، الوصول 06:45، الحد المرن للرحلة " + str(C.MAX_RIDE_MIN) + " د.",
    "حدود الباص المتوسط والفان والسرعات تقديرية وتحتاج معايرة ميدانية مع المشغّل.",
    "لا تُقاس: الارتفاعات، المطبات، الميول، السيارات الواقفة، الاتجاه الواحد في الشوارع المحلية، الإشارات، الزحمة.",
    "تنبيه خصوصية: هذا الملف فيه أسماء قُصّر، للاستخدام الداخلي فقط.",
]


def write_excel(path, df, by_school, by_district, pickups, fleet_sheet, quality, school_ops=None,
                scenarios=None, route_totals=None):
    wb = Workbook()
    wb.remove(wb.active)

    # ── الطلاب ──
    sd = pd.DataFrame({
        "الاسم": df["name"], "المدرسة": df["school"], "المرحلة": df["stage"], "الحي": df["district"],
        "المنطقة": df["zone"], "مستوى الباص الكبير": df["level_L"], "الدرجة": df["score_L"],
        "المشي (م)": df["walk_used_L"].round(0), "الثقة": df["conf_L"], "الأسباب": df["reasons_L"],
        "مستوى الباص المتوسط": df["level_M"], "مستوى الفان": df["level_S"],
        "المركبة الموصى بها": df["rec_vehicle_ar"],
        "رقم المسار": df.get("route", pd.Series([None] * len(df), index=df.index)),
        "وقت الركوب": df.get("board_time", pd.Series([None] * len(df), index=df.index)),
        "زمن الرحلة بالباص (د)": df.get("ride_min", pd.Series([np.nan] * len(df), index=df.index)).round(0),
        "زمن الطريق المباشر للمدرسة (د)": df.get("road_min", pd.Series([np.nan] * len(df), index=df.index)).round(0),
        "مسافة الطريق المباشر (كم)": df.get("road_km", pd.Series([np.nan] * len(df), index=df.index)).round(1),
    })
    ws_s = _sheet(wb, "الطلاب")
    _put_df(ws_s, sd, widths={"الأسباب": 70, "الاسم": 28, "المدرسة": 30})
    n = len(sd) + 1
    for name in ("مستوى الباص الكبير", "مستوى الباص المتوسط", "مستوى الفان"):
        _level_cf(ws_s, f"{_col(sd.columns, name)}2:{_col(sd.columns, name)}{n}")
    lvL, lvM, lvS = (_col(sd.columns, c) for c in ("مستوى الباص الكبير", "مستوى الباص المتوسط", "مستوى الفان"))
    sch, dist, rec = (_col(sd.columns, c) for c in ("المدرسة", "الحي", "المركبة الموصى بها"))

    # ── الملخص (معادلات) ──
    ws = _sheet(wb, "الملخص")
    wb.move_sheet("الملخص", offset=-1)
    ws["A1"] = "ملخص وصول الباص — جدة"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws.append([])
    ws.append(["المستوى", "العدد", "النسبة"])
    for lv in FILL:
        r = ws.max_row + 1
        ws.append([lv, f"=COUNTIF('الطلاب'!{lvL}:{lvL},A{r})", None])
    tot_row = ws.max_row + 1
    ws.append(["الإجمالي", f"=SUM(B4:B6)", "=1"])
    for r in range(4, 7):
        ws.cell(row=r, column=3).value = f"=B{r}/$B${tot_row}"
        ws.cell(row=r, column=3).number_format = "0%"
        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=FILL[ws.cell(row=r, column=1).value])
    ws.append([])
    ws.append(["الثقة المنخفضة", f"=COUNTIF('الطلاب'!{_col(sd.columns, 'الثقة')}:{_col(sd.columns, 'الثقة')},\"منخفضة\")"])
    ws.append(["عدد نقاط التجميع", len(pickups) if pickups is not None else 0])
    ws.append(["الإصدار", C.VERSION])
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name=FONT, bold=c.font.bold, size=c.font.size)
            c.alignment = Alignment(readingOrder=2)
    ws.column_dimensions["A"].width = 28

    # ── حسب المدرسة (COUNTIFS) ──
    schools = sorted(df["school"].unique())
    ws = _sheet(wb, "حسب المدرسة")
    hdr = ["المدرسة", "الطلاب", "سهل", "متوسط", "صعب", "% صعب"]
    ws.append(hdr)
    for i, s in enumerate(schools, 2):
        ws.append([s, f"=COUNTIF('الطلاب'!{sch}:{sch},A{i})"] +
                  [f"=COUNTIFS('الطلاب'!{sch}:{sch},$A{i},'الطلاب'!{lvL}:{lvL},\"{lv}\")" for lv in FILL] +
                  [f"=IF(B{i}=0,0,E{i}/B{i})"])
        ws.cell(row=i, column=6).number_format = "0%"
    _style_plain(ws, hdr, 36)

    # ── حسب الحي ──
    ds = sorted(df["district"].unique())
    ws = _sheet(wb, "حسب الحي")
    hdr = ["الحي", "الطلاب", "سهل", "متوسط", "صعب", "% صعب"]
    ws.append(hdr)
    for i, s in enumerate(ds, 2):
        ws.append([s, f"=COUNTIF('الطلاب'!{dist}:{dist},A{i})"] +
                  [f"=COUNTIFS('الطلاب'!{dist}:{dist},$A{i},'الطلاب'!{lvL}:{lvL},\"{lv}\")" for lv in FILL] +
                  [f"=IF(B{i}=0,0,E{i}/B{i})"])
        ws.cell(row=i, column=6).number_format = "0%"
    _style_plain(ws, hdr, 28)

    # ── مقارنة المركبات (COUNTIF) ──
    ws = _sheet(wb, "مقارنة المركبات")
    hdr = ["المركبة", "سهل", "متوسط", "صعب", "الموصى بها (عدد)"]
    ws.append(hdr)
    for name, col, rv in (("باص كبير", lvL, "باص كبير"), ("باص متوسط", lvM, "باص متوسط"), ("فان", lvS, "فان")):
        i = ws.max_row + 1
        ws.append([name] + [f"=COUNTIF('الطلاب'!{col}:{col},\"{lv}\")" for lv in FILL] +
                  [f"=COUNTIF('الطلاب'!{rec}:{rec},\"{rv}\")"])
    i = ws.max_row + 1
    ws.append(["نقطة تجميع / ترتيب خاص", None, None, None, f"=COUNTIF('الطلاب'!{rec}:{rec},\"نقطة تجميع / ترتيب خاص\")"])
    _style_plain(ws, hdr, 28)

    # ── نقاط التجميع ──
    ws = _sheet(wb, "نقاط التجميع")
    hdr = ["رقم", "عدد الطلاب", "الصعبين", "أبعد مشي (م)", "المدارس", "الحي", "الموقع"]
    ws.append(hdr)
    if pickups is not None and len(pickups):
        lon, lat = lonlat(pickups.x, pickups.y)
        for k, r in enumerate(pickups.itertuples(), 2):
            ws.append([r.pickup_id, r.students, r.hard, round(r.max_walk), r.schools, r.district,
                       f'=HYPERLINK("https://www.google.com/maps?q={lat[k - 2]:.6f},{lon[k - 2]:.6f}","خريطة")'])
    _style_plain(ws, hdr, 36)

    # ── الأسطول ──
    ws = _sheet(wb, "الأسطول")
    _put_df(ws, fleet_sheet)
    # ── جودة البيانات ──
    ws = _sheet(wb, "جودة البيانات")
    _put_df(ws, quality, widths={"الفحص": 44, "الإجراء_المقترح": 50})

    # ── التشغيل والمسارات ──
    if school_ops is not None and len(school_ops):
        so = school_ops.rename(columns={
            "school": "المدرسة", "need_large": "كبير مطلوب", "need_medium": "متوسط مطلوب", "need_small": "فان مطلوب",
            "need_total": "إجمالي المطلوب", "buses_large": "كبير حالي", "buses_medium": "متوسط حالي",
            "current_total": "إجمالي الحالي", "gap": "الفرق", "students": "الطلاب", "km_day": "كم/يوم",
            "hours_day": "ساعات/يوم", "fuel_sar_year": "وقود ريال/سنة", "co2_t_year": "CO2 طن/سنة",
            "avg_ride_min": "متوسط الرحلة (د)", "max_ride_min": "أطول رحلة (د)", "earliest_start": "أبكر انطلاق"})
        so = so.round(1)
        ws = _sheet(wb, "التشغيل والمسارات")
        _put_df(ws, so, widths={"المدرسة": 36})
        last = len(so) + 1
        tr = last + 1
        ws.cell(row=tr, column=1, value="الإجمالي").font = Font(name=FONT, bold=True)
        for j, c in enumerate(so.columns, 1):
            if c in ("كبير مطلوب", "متوسط مطلوب", "فان مطلوب", "إجمالي المطلوب", "كبير حالي", "متوسط حالي",
                     "إجمالي الحالي", "الفرق", "الطلاب", "كم/يوم", "ساعات/يوم", "وقود ريال/سنة", "CO2 طن/سنة"):
                ws.cell(row=tr, column=j, value=f"=SUM({CL(j)}2:{CL(j)}{last})").font = Font(name=FONT, bold=True)
    if scenarios is not None and len(scenarios):
        ws = _sheet(wb, "سيناريوهات الأسطول")
        _put_df(ws, scenarios.round(1))

    ws = _sheet(wb, "المنهجية")
    for t in METHOD:
        ws.append([t])
    ws.column_dimensions["A"].width = 120
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name=FONT, bold=(c.row == 1))
            c.alignment = Alignment(readingOrder=2, wrap_text=True)
    for w in wb.worksheets:
        w.sheet_view.rightToLeft = True
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def _style_plain(ws, hdr, first_w):
    for j in range(1, len(hdr) + 1):
        c = ws.cell(row=1, column=j)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HDR
        c.alignment = Alignment(horizontal="center", readingOrder=2, wrap_text=True)
        ws.column_dimensions[CL(j)].width = first_w if j == 1 else 16
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT)
            c.alignment = Alignment(readingOrder=2)
    ws.freeze_panes = "A2"
