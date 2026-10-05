import numpy as np
import pandas as pd
import pytest
from helpers import PROJ, X0, Y0, lines_xy, pt

from transport_hub.core import congestion as CG
from transport_hub.core import demo_city as DC
from transport_hub.core import geo
from transport_hub.core.access import Access, load_access
from transport_hub.core.roadnet import RoadNetwork
from transport_hub.core.store import Workspace
from transport_hub.ops import performance as P
from transport_hub.transit import gtfs, priority, service

TRUE = {"ذروة صباحية": 2.0, "منتصف النهار": 1.2, "ذروة مسائية": 1.8, "مساء": 1.0, "ليل": 1.0}
HOURS = {"ذروة صباحية": 7, "منتصف النهار": 12, "ذروة مسائية": 17, "مساء": 21, "ليل": 3}


def corridor(n=10, step=500.0):
    return RoadNetwork(lines_xy([[(i * step, 0) for i in range(n)]]), PROJ)


def make_feed_and_avl(net, slow_pair=None, days=3, per_period=4):
    """ممر مستقيم: محطات كل 1000 م. زمن القيادة الفعلي = زمن السير الحر × عامل الفترة الحقيقي (معروف)."""
    stop_x = [0, 1000, 2000, 3000, 4000]
    lo, la = PROJ.lonlat([X0 + x for x in stop_x], [Y0] * 5)
    stops = pd.DataFrame({"stop_id": list("ABCDE"), "stop_name": list("ABCDE"), "stop_lon": lo, "stop_lat": la})
    free_1000 = 1000 / (28 / 3.6)
    trips, st, avl = [], [], []
    for p, h in HOURS.items():
        for k in range(per_period):
            tid = f"t_{h}_{k}"
            trips.append(dict(route_id="X", service_id="WK", trip_id=tid, direction_id=0))
            t0 = h * 3600 + k * 600
            sched = t0
            act = t0
            for i, sid in enumerate("ABCDE"):
                st.append(dict(trip_id=tid, arrival_time=_hms(sched), departure_time=_hms(sched + 30), stop_id=sid, stop_sequence=i))
                for d in range(days):
                    avl.append(dict(date=f"2025-03-0{d + 1}", route_id="X", trip_id=tid, stop_id=sid, scheduled=sched, actual=int(act)))
                if i < 4:
                    f = TRUE[p] * (slow_pair[1] if slow_pair and slow_pair[0] == i and p == "ذروة صباحية" else 1.0)
                    sched += 30 + 150
                    act += 30 + free_1000 * f
    feed = gtfs.from_tables(dict(stops=stops, routes=pd.DataFrame({"route_id": ["X"]}), trips=pd.DataFrame(trips), stop_times=pd.DataFrame(st)))
    return feed, P.clean_avl(pd.DataFrame(avl))


def _hms(s):
    s = int(s)
    return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}"


def test_period_boundaries():
    assert [CG.period_of_hour(h) for h in (0, 5, 6, 8, 9, 14, 15, 18, 19, 23)] == [
        "ليل",
        "ليل",
        "ذروة صباحية",
        "ذروة صباحية",
        "منتصف النهار",
        "منتصف النهار",
        "ذروة مسائية",
        "ذروة مسائية",
        "مساء",
        "مساء",
    ]
    assert CG.period_of_seconds(25 * 3600) == "ليل"  # أوقات GTFS بعد منتصف الليل


def test_learns_known_factors_from_synthetic_avl():
    net = corridor()
    feed, avl = make_feed_and_avl(net)
    prof = CG.learn_from_avl(net, feed, avl, PROJ)
    for p, f in TRUE.items():
        assert prof.factor(p) == pytest.approx(f, rel=0.02), p
    assert prof.diagnostics.set_index("period").loc["ذروة صباحية", "pairs"] == 4
    assert prof.compatible(net) and prof.net_fp == net.fingerprint and prof.source.startswith("تعلّم")


def test_slow_segment_identified_as_worst_link():
    net = corridor()
    feed, avl = make_feed_and_avl(net, slow_pair=(2, 3.0))  # الزوج C→D أبطأ 3 أضعاف في الذروة الصباحية
    prof = CG.learn_from_avl(net, feed, avl, PROJ)
    w = CG.worst_links(prof, net, "ذروة صباحية", 2)
    mid_x = (w["xa"] + w["xb"]) / 2 - X0
    assert (mid_x > 2000 - 1).all() and (mid_x < 3000 + 1).all()  # أسوأ الأضلاع داخل C–D
    assert w["factor"].min() > 4.5 * 0 + prof.factor("ذروة صباحية") * 1.15
    # الأضلاع خارج المقطع البطيء أقرب للعامل العام
    f = prof.edge_factor["ذروة صباحية"]
    idx = np.where((net._du < net._dv) & ((net.node_xy[net._du, 0] - X0) < 1500))[0][0]
    assert f[idx] == pytest.approx(2.0, rel=0.05)


def test_global_profile_scales_times_by_period_in_both_modes():
    prof = CG.Profile({**{p: 1.0 for p in CG.PERIOD_KEYS}, "ذروة صباحية": 2.0}, source="اختبار")
    a, b = pt(0, 0), pt(2000, 0)
    est = Access(profile=prof)
    t_am, t_night = est.with_period("ذروة صباحية").drive_pairs(a, b, "time")[0], est.with_period("ليل").drive_pairs(a, b, "time")[0]
    assert t_am == pytest.approx(2 * t_night, rel=1e-6)
    assert est.drive_pairs(a, b, "time")[0] == pytest.approx(2000 * 1.3 / (37.5 / 3.6) * 1.25, rel=1e-6)  # بلا فترة: الافتراضي (30 كم/س)
    net = corridor()
    osm = Access(net, profile=prof)
    assert osm.with_period("ذروة صباحية").drive_pairs(a, b, "time")[0] == pytest.approx(
        2 * osm.with_period("ليل").drive_pairs(a, b, "time")[0], rel=0.01
    )
    assert osm.drive_pairs(a, b, "length")[0] == pytest.approx(2000, abs=2)  # المسافة لا تتأثر بالازدحام


def test_routing_avoids_congested_edges_only_in_peak():
    # مساران من A إلى B: مباشر 1000 م على y=0، والتفاف 1400 م عبر y=200
    segs = [[(0, 0), (500, 0), (1000, 0)], [(0, 0), (0, 200), (500, 200), (1000, 200), (1000, 0)]]
    net = RoadNetwork(lines_xy(segs), PROJ)
    my = (net.node_xy[net._du, 1] + net.node_xy[net._dv, 1]) / 2 - Y0
    direct = my < 1
    peak = np.where(direct, 4.0, 1.25)
    free = np.full(net.n_edges, 1.25)
    prof = CG.Profile({p: 1.25 for p in CG.PERIOD_KEYS}, {"ذروة صباحية": peak, "ليل": free}, np.ones(net.n_edges), net.fingerprint, "اختبار")
    a, b = pt(0, 0), pt(1000, 0)
    p_peak = net.drive_path(a, b, "time", prof, "ذروة صباحية")
    p_night = net.drive_path(a, b, "time", prof, "ليل")
    assert len(p_peak) > len(p_night) and np.ptp(p_peak[:, 1]) > 100  # في الذروة يلتف عبر y=200
    assert np.ptp(p_night[:, 1]) < 1  # ليلاً المباشر
    assert net.drive_pairs(a, b, "time", profile=prof, period="ذروة صباحية")[0] < 4 * 1000 / (28 / 3.6)  # الالتفاف أرخص من البقاء في الازدحام


def test_manual_hourly_curve():
    df = pd.DataFrame({"hour": [7, 8, 13, 17, 18, 22], "factor": [1.6, 1.8, 1.1, 1.7, 1.5, 1.0]})
    p = CG.from_hourly(df)
    assert (
        p.factor("ذروة صباحية") == pytest.approx(1.7)
        and p.factor("منتصف النهار") == pytest.approx(1.1)
        and p.factor("ذروة مسائية") == pytest.approx(1.6)
    )
    assert p.edge_factor is None and p.compatible(corridor())
    p2 = CG.from_hourly(pd.DataFrame({"hour": [8], "speed_ratio": [0.5]}))
    assert p2.factor("ذروة صباحية") == pytest.approx(2.0)
    with pytest.raises(ValueError, match="hour"):
        CG.from_hourly(pd.DataFrame({"a": [1]}))


def test_incompatible_profile_is_ignored(tmp_path):
    w = Workspace(tmp_path)
    w.save_obj("proj_epsg", 32637)
    w.save_obj("roads_lines", lines_xy([[(i * 500, 0) for i in range(10)]]))
    net = load_access(w).net
    good = CG.Profile(
        {p: 1.5 for p in CG.PERIOD_KEYS}, {p: np.full(net.n_edges, 1.5) for p in CG.PERIOD_KEYS}, np.ones(net.n_edges), net.fingerprint, "تعلّم"
    )
    w.save_obj("congestion", good)
    assert load_access(w).profile is not None
    assert load_access(w, use_congestion=False).profile is None
    w.save_obj(
        "congestion",
        CG.Profile(
            {p: 1.5 for p in CG.PERIOD_KEYS}, {p: np.full(net.n_edges, 1.5) for p in CG.PERIOD_KEYS}, np.ones(net.n_edges), "شبكة أخرى", "تعلّم"
        ),
    )
    assert load_access(w).profile is None  # من شبكة مختلفة: لا يُطبَّق
    w2 = Workspace(tmp_path / "w2")  # بدون شوارع: منحنى عام فقط
    w2.save_obj("congestion", CG.default_profile(1.4))
    a = load_access(w2)
    assert a.mode == "estimate" and a.profile is not None


def test_learning_errors_are_clear():
    net = corridor()
    feed, avl = make_feed_and_avl(net)
    with pytest.raises(ValueError, match="لا ملاحظات كافية"):
        CG.learn_from_avl(net, feed, avl.iloc[:0], PROJ)
    other = avl.assign(trip_id="غير موجود")
    with pytest.raises(ValueError, match="لا ملاحظات كافية"):
        CG.learn_from_avl(net, feed, other, PROJ)


def test_route_time_check_detects_unrealistic_schedule():
    net = corridor()
    feed, avl = make_feed_and_avl(net)
    prof = CG.learn_from_avl(net, feed, avl, PROJ)
    acc = Access(net, profile=prof)
    tab = service.route_time_check(feed, PROJ, acc)
    am = tab[tab["period"] == "ذروة صباحية"].iloc[0]
    night = tab[tab["period"] == "ليل"].iloc[0]
    assert am["predicted_min"] > am["scheduled_min"] and am["verdict"].startswith("ناقص")  # الجدول بُني على 150 ث/مقطع والواقع أبطأ
    assert am["predicted_min"] > night["predicted_min"]  # الذروة أطول من الليل
    assert set(tab["period"]) == set(CG.PERIOD_KEYS) and tab["delta_min"].round(1).equals(tab["delta_min"])
    # جدول واقعي: نعيد توقيت رحلات الذروة الصباحية بزمن القيادة المتوقع بين المحطات (مع بقاء التوقف 30 ث) فيصير الحكم «واقعي»
    trips_am = feed.trips[feed.trips["trip_id"].str.startswith("t_7_")]
    stx = feed.stops.set_index("stop_id")
    xy = np.column_stack(PROJ.xy(stx.loc[list("ABCDE"), "stop_lon"], stx.loc[list("ABCDE"), "stop_lat"]))
    drive = acc.with_period("ذروة صباحية").drive_pairs(xy[:-1], xy[1:], "time")
    st = feed.stop_times[feed.stop_times["trip_id"].isin(trips_am["trip_id"])].copy().sort_values(["trip_id", "stop_sequence"])
    for _tid, g in st.groupby("trip_id"):
        t = g["dep"].iloc[0] - 30
        for j, ix in enumerate(g.index):
            st.loc[ix, "arr"], st.loc[ix, "dep"] = t, t + 30
            t = t + 30 + (drive[j] if j < len(drive) else 0)
    feed2 = gtfs.Feed(feed.stops, feed.routes, trips_am, st, feed.calendar)
    ok = service.route_time_check(feed2, PROJ, acc)
    assert ok.iloc[0]["verdict"] == "واقعي" and abs(ok.iloc[0]["delta_min"]) < 0.06 * ok.iloc[0]["scheduled_min"]


def test_bus_priority_savings_logic():
    net = corridor()
    feed, avl = make_feed_and_avl(net, slow_pair=(2, 3.0))
    prof = CG.learn_from_avl(net, feed, avl, PROJ)
    acc = Access(net, profile=prof)
    r0 = priority.evaluate(feed, PROJ, acc, "X", lane_share=0.4, delay_cut=0.0)
    assert (r0["table"]["saved_min"] == 0).all()  # بلا تقليل تأخير لا وفر
    a = priority.evaluate(feed, PROJ, acc, "X", lane_share=0.2, delay_cut=0.6, pax_day=1000)
    b = priority.evaluate(feed, PROJ, acc, "X", lane_share=0.8, delay_cut=0.6, pax_day=1000)
    assert (b["table"]["saved_min"].to_numpy() >= a["table"]["saved_min"].to_numpy() - 1e-9).all() and b["table"]["saved_min"].sum() > a["table"][
        "saved_min"
    ].sum() > 0
    c = priority.evaluate(feed, PROJ, acc, "X", lane_share=0.8, delay_cut=0.9)
    assert c["table"]["saved_min"].sum() > b["table"]["saved_min"].sum()
    am = b["table"].set_index("period").loc["ذروة صباحية"]
    ng = b["table"].set_index("period").loc["ليل"]
    assert am["saved_min"] > ng["saved_min"]  # الوفر في الذروة أكبر
    assert b["summary"]["vehicle_hours_saved_year"] > 0 and b["summary"]["passenger_value_year"] > 0
    assert b["links"] is not None and len(b["links"]) > 0
    est = priority.evaluate(feed, PROJ, Access(profile=prof), "X")  # بدون شبكة: يعمل بعامل عام
    assert est["table"]["saved_min"].sum() > 0
    with pytest.raises(ValueError, match="غير موجود"):
        priority.evaluate(feed, PROJ, acc, "ZZZ")


def test_deadhead_time_follows_period_congestion():
    from transport_hub.ops import blocking as B

    net = corridor()
    prof = CG.Profile({**{p: 1.0 for p in CG.PERIOD_KEYS}, "ذروة صباحية": 3.0}, source="اختبار")
    acc = Access(net, profile=prof)
    base = dict(route_id="R", direction_id=0, from_x=X0, from_y=Y0, to_x=X0 + 2000, to_y=Y0, km=1)
    t = pd.DataFrame([dict(trip_id="am", start=6.5 * 3600, end=7.5 * 3600, **base), dict(trip_id="am2", start=9 * 3600, end=9.5 * 3600, **base)])
    t.loc[:, ["from_x", "to_x"]] = [[X0, X0 + 2000], [X0, X0 + 2000]]
    _, ds_am = B._deadhead(t.assign(end=[7.5 * 3600, 7.5 * 3600]), acc)
    _, ds_night = B._deadhead(t.assign(end=[2 * 3600, 2 * 3600]), acc)
    assert ds_am[0, 0] == pytest.approx(3 * ds_night[0, 0], rel=0.02)


def test_outputs_include_congestion_when_profile_exists(tmp_path):
    import io
    import json
    import zipfile

    from transport_hub.core import data as D
    from transport_hub.exports import builder

    w = Workspace(tmp_path)
    d = DC.build_all()
    p = geo.Projector.for_points(d["population"].lon, d["population"].lat)
    w.save_obj("proj_epsg", p.epsg)
    w.save_df("population", D.clean_population(d["population"], p))
    w.save_obj("gtfs", d["gtfs"])
    w.save_obj("roads_lines", DC.street_lines())
    avl, _ = DC.avl_apc(d["gtfs"], days=2)
    w.save_df("avl", avl)
    DC.learn_demo_congestion(w)
    assert w.obj("congestion") is not None
    z, meta = builder.build_package(w)
    zf = zipfile.ZipFile(io.BytesIO(z))
    names = zf.namelist()
    assert "gis/أضلاع_مزدحمة.geojson" in names
    gj = json.loads(zf.read("gis/أضلاع_مزدحمة.geojson"))
    assert gj["features"] and gj["features"][0]["geometry"]["type"] == "LineString"
    assert "الازدحام" in meta["params"] and "×" in meta["params"]["الازدحام"]
    m2, sheets2, _, _ = builder.assemble(w, use_congestion=False)
    assert m2["params"]["الازدحام"].startswith("غير مطبّق") and m2["fingerprint"] != meta["fingerprint"]
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(zf.read("التقرير.xlsx")))
    assert "واقعية الجدول" in wb.sheetnames and "الازدحام - الفترات" in wb.sheetnames


def test_api_congestion_endpoints(tmp_path, monkeypatch):
    import importlib
    import json as _json

    from fastapi.testclient import TestClient

    monkeypatch.setenv("TRANSPORT_HUB_WORKSPACE", str(tmp_path))
    monkeypatch.setenv("TRANSPORT_HUB_API_KEYS", _json.dumps({"a": "admin"}))
    import transport_hub.core.store as store

    importlib.reload(store)
    import transport_hub.api as api

    importlib.reload(api)
    c = TestClient(api.app)
    h = {"X-API-Key": "a"}
    assert c.get("/congestion/profile", headers=h).status_code == 409
    c.post("/demo/load", headers=h)
    prof = c.get("/congestion/profile", headers=h).json()
    assert len(prof["periods"]) == 5 and "تعلّم" in prof["source"]
    rtc = c.get("/transit/route-time-check", headers=h).json()
    assert rtc and {"period", "predicted_min", "verdict"} <= set(rtc[0])
