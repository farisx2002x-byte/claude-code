import io
import json
import zipfile

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from transport_hub.core import data as D
from transport_hub.core import demo_city, geo, safety
from transport_hub.core.store import Workspace
from transport_hub.exports import builder, html_report, meta
from transport_hub.exports import dictionary as DICT
from transport_hub.exports import excel as XL
from transport_hub.exports import geo as XG


@pytest.fixture(scope="module")
def ws(tmp_path_factory):
    w = Workspace(tmp_path_factory.mktemp("ws"))
    d = demo_city.build_all()
    p = geo.Projector.for_points(d["population"].lon, d["population"].lat)
    w.save_obj("proj_epsg", p.epsg)
    w.save_df("population", D.clean_population(d["population"], p))
    w.save_df("poi", D.clean_poi(d["poi"], p))
    w.save_obj("gtfs", d["gtfs"])
    w.save_df("trips", D.clean_trips(d["trips"], p))
    w.save_df("stands", p.attach(d["stands"]))
    avl, _ = demo_city.avl_apc(d["gtfs"], days=2)
    w.save_df("avl", avl)
    for k in ("population", "poi", "gtfs", "trips", "stands", "avl"):
        w.set_source(k, "demo")
    return w


def test_package_contents_and_fingerprint(ws):
    z, m = builder.build_package(ws)
    zf = zipfile.ZipFile(io.BytesIO(z))
    names = set(zf.namelist())
    assert {"التقرير.xlsx", "التقرير_التنفيذي.html", "manifest.json", "README.txt"} <= names
    assert any(n.startswith("gis/") and n.endswith(".geojson") for n in names)
    manifest = json.loads(zf.read("manifest.json"))
    assert manifest["fingerprint"] == m["fingerprint"] and manifest["demo"] is True
    assert m["fingerprint"] in zf.read("README.txt").decode()
    assert "بيانات تجريبية" in zf.read("التقرير_التنفيذي.html").decode()
    # الإنتاج حتمي: نفس البيانات والمعاملات = نفس البصمة
    assert builder.assemble(ws)[0]["fingerprint"] == m["fingerprint"]
    assert builder.assemble(ws, radius=500)[0]["fingerprint"] != m["fingerprint"]


def test_excel_structure_and_formats(ws):
    m, sheets, _, _ = builder.assemble(ws)
    buf = io.BytesIO()
    XL.build_workbook(buf, sheets, m)
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    assert wb.sheetnames[0] == "الغلاف" and wb.sheetnames[-1] == "قاموس البيانات"
    assert all(len(n) <= 31 for n in wb.sheetnames)
    for w in wb.worksheets:
        assert w.sheet_view.rightToLeft
    ws1 = wb["النقل العام - الخطوط"]
    hdr_row = next(r for r in range(1, 6) if ws1.cell(row=r, column=1).value == "الخط")
    assert ws1.freeze_panes == f"A{hdr_row + 1}" and ws1.auto_filter.ref
    # رؤوس عربية وأنماط أرقام من القاموس
    heads = [c.value for c in ws1[hdr_row]]
    assert "أسطول الذروة" in heads and "السرعة التجارية" in heads
    j = heads.index("الطول") + 1
    assert ws1.cell(row=hdr_row + 1, column=j).number_format == "#,##0.0"
    # نص بأرقام (06:00–22:00) يقرأ LTR
    k = heads.index("ساعات الخدمة") + 1
    assert ws1.cell(row=hdr_row + 1, column=k).alignment.readingOrder == 1
    assert len(ws1._charts) == 1
    # لا قيم nan/inf في أي خلية
    for w in wb.worksheets:
        for row in w.iter_rows(values_only=True):
            assert not any(isinstance(v, float) and (np.isnan(v) or np.isinf(v)) for v in row)
    cover = wb["الغلاف"]
    texts = [c.value for r in cover.iter_rows() for c in r if c.value]
    assert m["fingerprint"] in texts and any("تجريبية" in str(t) for t in texts)
    assert any(c.hyperlink for r in cover.iter_rows() for c in r)


def test_dictionary_consistency():
    for col, (lab, _desc, _unit, fmt) in DICT.COLUMNS.items():
        assert lab and fmt in XL.FORMATS, col
    df = pd.DataFrame({"route_id": ["R1"], "length_km": [1.0], "unknown": [1]})
    r = DICT.relabel(df)
    assert list(r.columns) == ["الخط", "الطول", "unknown"]


def test_html_report_sanity(ws):
    m, _, sections, _ = builder.assemble(ws)
    h = html_report.build(m, sections)
    assert h.startswith("<!doctype html>") and 'dir="rtl"' in h and "<svg" in h
    assert "<script" not in h.lower()  # ملف ساكن بلا سكربتات
    assert "&lt;" not in h or "<script" not in h
    # الأوقات داخل bdi LTR كي لا تنعكس
    assert '<bdi dir="ltr">06:00–22:00</bdi>' in h
    assert html_report.nice_max(9) == 10 and html_report.nice_max(0) == 1.0


def test_html_escapes_untrusted_text():
    m = meta.build("<script>alert(1)</script>", {"k": "<b>x</b>"})
    h = html_report.build(m, [dict(title="<img src=x onerror=1>", table=pd.DataFrame({"name": ["<script>x</script>"]}))])
    assert "<script>alert" not in h and "<img src=x" not in h and "&lt;script&gt;" in h


def test_geo_exports(ws):
    df = pd.DataFrame({"rank": [1, 2], "gain": [10.5, np.nan], "lon": [39.1, 39.2], "lat": [21.5, 21.6]})
    gj = json.loads(XG.to_geojson(df))
    assert len(gj["features"]) == 2 and gj["features"][1]["properties"]["gain"] is None
    assert gj["features"][0]["geometry"]["coordinates"] == [39.1, 21.5]
    kml = XG.to_kml(df, doc_name="اختبار").decode()
    assert "<Placemark>" in kml and kml.count("<Placemark>") == 2
    assert XG.to_csv(df).startswith(b"\xef\xbb\xbf")  # BOM لفتح العربي في Excel


def test_missing_data_graceful(tmp_path):
    w = Workspace(tmp_path)
    with pytest.raises(ValueError, match="لا بيانات كافية"):
        builder.build_package(w)
    d = demo_city.build_all()
    p = geo.Projector.for_points(d["population"].lon, d["population"].lat)
    w.save_obj("proj_epsg", p.epsg)
    w.save_df("trips", D.clean_trips(d["trips"], p))
    m, sheets, _, _ = builder.assemble(w)  # التاكسي وحده يكفي
    assert any("التاكسي" in s.name for s in sheets) and any("النقل العام" in n for n in m["notes"])
    assert not m["demo"]


# ───────── أمان المدخلات وتنظيف البيانات ─────────
def test_safety_checks():
    with pytest.raises(safety.UnsafeUpload):
        safety.check_size(2 * 10**9)
    with pytest.raises(safety.UnsafeUpload, match="zip"):
        safety.check_zip(b"not a zip")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("../evil.txt", "x")
    with pytest.raises(safety.UnsafeUpload, match="غير آمن"):
        safety.check_zip(buf.getvalue())
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("a.txt", b"0" * 50_000_000)
    with pytest.raises(safety.UnsafeUpload, match="ضغط"):
        safety.check_zip(bomb.getvalue())
    ok = io.BytesIO()
    with zipfile.ZipFile(ok, "w") as z:
        z.writestr("stops.txt", "stop_id\n1\n")
    safety.check_zip(ok.getvalue())


def test_cleaners_report_issues():
    proj = geo.Projector(32637)
    raw = pd.DataFrame(
        {"zone_id": ["a", "b", "b", "c", "d"], "pop": [100, -5, 20, "x", 50], "lon": [39.1, 39.1, 39.1, 39.1, 0], "lat": [21.5, 21.5, 21.5, 21.5, 0]}
    )
    out = D.clean_population(raw, proj)
    assert list(out["zone_id"]) == ["a", "b"] or list(out["zone_id"]) == ["a"]
    iss = out.attrs["issues"]
    assert iss["سكان سالب"] == 1 and iss["سكان فارغ/غير رقمي"] == 1 and iss["إحداثيات غير صالحة"] == 1 and iss["معرّف مكرر"] == 1
    assert out.attrs["rows_in"] == 5 and out.attrs["rows_out"] == len(out)
    rep = D.quality_report(pop=out, load_report={"population": out.attrs})
    assert (rep["الفحص"].str.contains("استُبعد")).any()
    with pytest.raises(D.DataError, match="أعمدة ناقصة"):
        D.clean_population(pd.DataFrame({"a": [1]}), proj)
    bad_trips = pd.DataFrame(
        {
            "pickup_time": ["2025-01-01 10:00", "bad", "2025-01-01 11:00"],
            "pickup_lon": [39.1, 39.1, 39.1],
            "pickup_lat": [21.5, 21.5, 21.5],
            "fare": [10, 5, -3],
        }
    )
    t = D.clean_trips(bad_trips, proj)
    assert len(t) == 2 and t.attrs["issues"]["وقت غير صالح"] == 1 and t.attrs["issues"]["أجرة سالبة"] == 1


def test_workspace_sources_and_demo_flag(tmp_path):
    w = Workspace(tmp_path)
    assert not w.is_demo()
    w.set_source("population", "upload")
    assert not w.is_demo()
    w.set_source("poi", "demo")
    assert w.is_demo()
    w.clear()
    assert not w.is_demo()
