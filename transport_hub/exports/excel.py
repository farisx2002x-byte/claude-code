"""مصنف Excel احترافي: غلاف وفهرس بروابط، أوراق منسّقة (RTL، أنماط أرقام، تصفية، تجميد، طباعة)، رسوم أصلية، وقاموس بيانات."""

from dataclasses import dataclass, field

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

from transport_hub.exports import dictionary as DICT

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
ZEBRA = PatternFill("solid", fgColor="F3F6FA")
THIN = Side(style="thin", color="D0D7DE")
FORMATS = {
    "int": "#,##0",
    "dec1": "#,##0.0",
    "dec2": "#,##0.00",
    "dec5": "0.00000",
    "pct": '0.0"%"',
    "ratio": "0.0%",
    "text": "@",
    "time": "hh:mm",
    "flag": "General",
}
STATUS_FILL = {
    "🟢": "C6EFCE",
    "🟡": "FFEB9C",
    "🔴": "FFC7CE",
    "سهل": "C6EFCE",
    "متوسط": "FFEB9C",
    "صعب": "FFC7CE",
    "متأخرة": "FFC7CE",
    "قريبة": "FFEB9C",
    "سليمة": "C6EFCE",
}


@dataclass
class Sheet:
    name: str
    df: pd.DataFrame
    title: str = ""
    notes: list = field(default_factory=list)
    chart: dict | None = None  # {"type": "bar"|"line", "x": col, "y": [cols], "title": str}
    status_cols: list = field(default_factory=list)


def _style_cell(c, bold=False, color=None):
    c.font = Font(name=FONT, bold=bold, color=color)
    v = c.value
    # نص بأرقام وبدون حروف عربية (مثل 06:00–22:00) يُقرأ LTR كي لا ينعكس داخل ورقة RTL
    ltr = isinstance(v, str) and any(ch.isdigit() for ch in v) and not any("\u0600" <= ch <= "\u06ff" for ch in v)
    c.alignment = Alignment(readingOrder=1 if ltr else 2, horizontal="right" if ltr else None, vertical="center", wrap_text=False)


def _val(v):
    if isinstance(v, np.bool_):
        return "نعم" if v else "لا"
    if isinstance(v, bool):
        return "نعم" if v else "لا"
    if isinstance(v, (np.floating, float)):
        return None if (np.isnan(v) or np.isinf(v)) else float(v)
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, (pd.Timestamp,)):
        return v.to_pydatetime()
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return None
    if isinstance(v, (int, str)):
        return v
    return str(v)


def _width(label, series):
    sample = [len(str(x)) for x in series.head(200) if x is not None]
    return min(max([len(str(label))] + sample[:200] or [8]) * 1.15 + 3, 48)


def _write_sheet(wb, sh, raw_cols):
    ws = wb.create_sheet(sh.name[:31])
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = "1F3864"
    row0 = 1
    if sh.title:
        ws.cell(row=1, column=1, value=sh.title).font = Font(name=FONT, bold=True, size=13)
        row0 = 2
        for n in sh.notes:
            ws.cell(row=row0, column=1, value=n).font = Font(name=FONT, italic=True, color="666666")
            row0 += 1
        row0 += 1
    df = sh.df.copy()
    keys = list(df.columns)
    for j, k in enumerate(keys, 1):
        c = ws.cell(row=row0, column=j, value=DICT.label(k))
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
        u = DICT.unit_of(k)
        if u and u != "عدد":
            c.comment = None
    ws.row_dimensions[row0].height = 30
    for i, rec in enumerate(df.itertuples(index=False), row0 + 1):
        for j, (k, v) in enumerate(zip(keys, rec), 1):
            c = ws.cell(row=i, column=j, value=_val(v))
            _style_cell(c)
            c.number_format = FORMATS.get(DICT.fmt_of(k), "General")
            c.border = Border(bottom=THIN)
            if (i - row0) % 2 == 0:
                c.fill = ZEBRA
    for j, k in enumerate(keys, 1):
        ws.column_dimensions[CL(j)].width = _width(DICT.label(k), df[k].map(lambda x: x if not isinstance(x, float) else round(x, 1)))
    last = row0 + len(df)
    if len(df):
        ws.auto_filter.ref = f"A{row0}:{CL(len(keys))}{last}"
    ws.freeze_panes = ws.cell(row=row0 + 1, column=1)
    for col in sh.status_cols:
        if col in keys:
            j = keys.index(col) + 1
            for val, color in STATUS_FILL.items():
                ws.conditional_formatting.add(
                    f"{CL(j)}{row0 + 1}:{CL(j)}{last}",
                    CellIsRule(operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", bgColor=color, fgColor=color)),
                )
    # طباعة
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{row0}:{row0}"
    ws.oddFooter.center.text = "&P / &N"
    # رسم أصلي
    if sh.chart and len(df):
        spec = sh.chart
        ch = LineChart() if spec.get("type") == "line" else BarChart()
        ch.title = spec.get("title", "")
        ch.height, ch.width = 8.5, 20
        xj = keys.index(spec["x"]) + 1
        for y in spec["y"]:
            yj = keys.index(y) + 1
            ch.add_data(Reference(ws, min_col=yj, min_row=row0, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=xj, min_row=row0 + 1, max_row=last))
        ws.add_chart(ch, f"{CL(len(keys) + 2)}{row0}")
    raw_cols.extend(keys)
    return ws, row0


def build_workbook(path, sheets, meta, with_dictionary=True):
    """path: مسار أو buffer. sheets: قائمة Sheet. meta: ناتج exports.meta.build."""
    wb = Workbook()
    cover = wb.active
    cover.title = "الغلاف"
    cover.sheet_view.rightToLeft = True
    cover.sheet_properties.tabColor = "C00000" if meta.get("demo") else "2E7D32"
    cover["A1"] = meta["title"]
    cover["A1"].font = Font(name=FONT, bold=True, size=18)
    info = [("الإصدار", meta["version"]), ("تاريخ الإنشاء", meta["generated"]), ("بصمة التشغيل", meta["fingerprint"])]
    r = 3
    for k, v in info:
        cover.cell(row=r, column=1, value=k).font = Font(name=FONT, bold=True)
        cover.cell(row=r, column=2, value=v).font = Font(name=FONT)
        r += 1
    if meta.get("params"):
        r += 1
        cover.cell(row=r, column=1, value="المعاملات المستخدمة").font = Font(name=FONT, bold=True, size=12)
        r += 1
        for k, v in meta["params"].items():
            cover.cell(row=r, column=1, value=str(k)).font = Font(name=FONT)
            cover.cell(row=r, column=2, value=_val(v) if not isinstance(v, (list, dict)) else str(v)).font = Font(name=FONT)
            r += 1
    if meta.get("datasets"):
        r += 1
        cover.cell(row=r, column=1, value="مصادر البيانات").font = Font(name=FONT, bold=True, size=12)
        r += 1
        for k, v in meta["datasets"].items():
            cover.cell(row=r, column=1, value=k).font = Font(name=FONT)
            cover.cell(row=r, column=2, value=f"{v.get('rows', '')} سجل — {v.get('loaded', '')}").font = Font(name=FONT)
            r += 1
    r += 1
    cover.cell(row=r, column=1, value="الفهرس").font = Font(name=FONT, bold=True, size=12)
    r += 1
    toc_row = r
    raw = []
    for sh in sheets:
        ws, _ = _write_sheet(wb, sh, raw)
        c = cover.cell(row=r, column=1, value=sh.name)
        c.hyperlink = f"#'{ws.title}'!A1"
        c.font = Font(name=FONT, color="0563C1", underline="single")
        cover.cell(row=r, column=2, value=f"{len(sh.df):,} صف").font = Font(name=FONT, color="666666")
        r += 1
    r += 1
    cover.cell(row=r, column=1, value="تنبيهات").font = Font(name=FONT, bold=True, size=12, color="C00000")
    for d in meta.get("disclaimers", []) + meta.get("notes", []):
        r += 1
        cover.cell(row=r, column=1, value="• " + d).font = Font(name=FONT)
    cover.column_dimensions["A"].width = 46
    cover.column_dimensions["B"].width = 60
    for row in cover.iter_rows():
        for c in row:
            c.alignment = Alignment(readingOrder=2, vertical="center")
    if with_dictionary:
        dic = DICT.dictionary_frame(raw)
        if dic:
            ws = wb.create_sheet("قاموس البيانات")
            ws.sheet_view.rightToLeft = True
            for j, h in enumerate(("العمود", "الوصف", "الوحدة"), 1):
                c = ws.cell(row=1, column=j, value=h)
                c.font, c.fill = Font(name=FONT, bold=True, color="FFFFFF"), HDR_FILL
            for i, rec in enumerate(dic, 2):
                for j, k in enumerate(("العمود", "الوصف", "الوحدة"), 1):
                    _style_cell(ws.cell(row=i, column=j, value=rec[k]))
            for col, w in zip("ABC", (28, 60, 12)):
                ws.column_dimensions[col].width = w
            c = cover.cell(row=toc_row + len(sheets), column=1, value="قاموس البيانات")
            c.hyperlink = "#'قاموس البيانات'!A1"
            c.font = Font(name=FONT, color="0563C1", underline="single")
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
