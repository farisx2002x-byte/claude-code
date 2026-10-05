"""تقارير: Excel عربي (RTL) من أي مجموعة جداول، وملخص تنفيذي HTML."""

import html
from datetime import datetime

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def excel(path, sheets, title="تقرير منصة النقل"):
    """sheets: dict اسم الورقة → DataFrame. كل الأوراق RTL بخط Arial."""
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "الغلاف"
    ws0.sheet_view.rightToLeft = True
    ws0["A1"], ws0["A2"] = title, f"أُنشئ: {datetime.now():%Y-%m-%d %H:%M}"
    ws0["A1"].font = Font(name="Arial", bold=True, size=16)
    for i, name in enumerate(sheets, 4):
        ws0.cell(row=i, column=1, value=name)
    ws0.column_dimensions["A"].width = 50
    for name, df in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.sheet_view.rightToLeft = True
        for j, c in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=j, value=str(c))
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3864")
            cell.alignment = Alignment(horizontal="center", readingOrder=2, wrap_text=True)
            ws.column_dimensions[get_column_letter(j)].width = min(max(len(str(c)) + 4, 12), 40)
        for i, row in enumerate(df.itertuples(index=False), 2):
            for j, v in enumerate(row, 1):
                if isinstance(v, (np.floating, float)):
                    v = None if (np.isnan(v) or np.isinf(v)) else float(v)
                elif isinstance(v, np.integer):
                    v = int(v)
                elif isinstance(v, (np.bool_,)):
                    v = bool(v)
                elif not isinstance(v, (int, str, bool, type(None))):
                    v = str(v)
                c = ws.cell(row=i, column=j, value=v)
                c.font = Font(name="Arial")
                c.alignment = Alignment(readingOrder=2)
        ws.freeze_panes = "A2"
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def executive_html(title, kpis, notes=None):
    """ملخص تنفيذي مستقل (HTML واحد، RTL، يدعم الوضع الداكن). kpis: DataFrame من scorecard.build."""
    rows = "".join(
        f"<tr><td>{html.escape(str(r.المؤشر))}</td><td>{r.الفعلي if r.الفعلي is not None else '—'} {html.escape(str(r.الوحدة))}</td>"
        f"<td>{r.المستهدف} {html.escape(str(r.الوحدة))}</td><td>{r.الحالة}</td></tr>"
        for r in kpis.itertuples()
    )
    notes_html = "".join(f"<li>{html.escape(n)}</li>" for n in (notes or []))
    return f"""<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8"><title>{html.escape(title)}</title>
<style>:root{{--bg:#fff;--fg:#1b1f23;--bd:#d0d7de}}@media(prefers-color-scheme:dark){{:root{{--bg:#111418;--fg:#e6e8eb;--bd:#2f3740}}}}
body{{font-family:Arial,Tahoma,sans-serif;background:var(--bg);color:var(--fg);max-width:900px;margin:24px auto;padding:0 16px}}
table{{width:100%;border-collapse:collapse}}td,th{{border-bottom:1px solid var(--bd);padding:8px;text-align:right}}</style></head>
<body><h1>{html.escape(title)}</h1><p>{datetime.now():%Y-%m-%d}</p>
<table><tr><th>المؤشر</th><th>الفعلي</th><th>المستهدف</th><th>الحالة</th></tr>{rows}</table><ul>{notes_html}</ul></body></html>"""
